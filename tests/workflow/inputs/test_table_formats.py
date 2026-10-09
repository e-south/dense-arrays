"""Optional table inputs preserve typed values and the shared part contract.

Author: Eric J. South.
"""

import builtins
import json
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from typer.testing import CliRunner

import dense_arrays as da
from dense_arrays import parts, planning
from dense_arrays.cli import app
from dense_arrays.parts.ingestion import read_parts
from dense_arrays.parts.serialization import table_from_dict, table_to_dict
from dense_arrays.workflow.inputs import read_source


def write_table(path: Path, rows: list[dict]) -> None:
    if path.suffix == ".parquet":
        pq.write_table(pa.Table.from_pylist(rows), path, row_group_size=1)
    else:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "parts"
        sheet.append(list(rows[0]))
        for row in rows:
            sheet.append(list(row.values()))
        book.save(path)
        book.close()


@pytest.mark.parametrize("format_name", ["parquet", "xlsx"])
def test_typed_tables_preserve_coordinates_metadata_and_normalization(
    tmp_path: Path, format_name: str
):
    path = tmp_path / f"parts.{format_name}"
    write_table(
        path,
        [
            {
                "id": "p1",
                "dna": " acgt ",
                "core_start": 0,
                "core_end": 2,
                "core_orientation": "forward",
                "quality": 1.5,
                "verified": True,
                "unused": "ignored",
            }
        ],
    )
    source = parts.PartTable(
        path,
        format_name,
        columns={"part_id": "id", "sequence": "dna"},
        metadata_columns={"quality": "quality", "verified": "verified"},
        normalization=parts.Normalization(uppercase=True, trim_outer_whitespace=True),
    )
    imported = read_parts(source)
    assert imported.parts == (
        parts.Part(
            "p1",
            "ACGT",
            core_start=0,
            core_end=2,
            core_orientation="forward",
            metadata={"quality": 1.5, "verified": True},
        ),
    )
    assert imported.ignored_columns == ("unused",)
    assert imported.report.transformations[0].before == " acgt "
    assert table_from_dict(table_to_dict(source)) == source


@pytest.mark.parametrize("format_name", ["parquet", "xlsx"])
def test_table_prepare_plan_and_cli_preserve_portable_parts(
    tmp_path: Path, format_name: str
):
    path = tmp_path / f"parts.{format_name}"
    write_table(
        path,
        [{"part_id": "p1", "sequence": "AAA"}, {"part_id": "p2", "sequence": "CCC"}],
    )
    recipe = parts.PreparationSpec(parts.PartTable(path, format_name))
    saved = tmp_path / "prepare.json"
    da.export(recipe, view="request", out=saved)
    python_pool = da.prepare(recipe, out=tmp_path / "python")
    result = CliRunner().invoke(
        app, ["prepare", str(saved), "--out", str(tmp_path / "cli"), "--json"]
    )
    assert result.exit_code == 0, result.output
    assert json.loads(result.stdout)["pool_id"] == python_pool.pool_id
    plan = da.plan(
        planning.DesignSpec(
            parts.PartTable(path, format_name),
            planning.Length(maximum=6),
            strands="single",
        )
    )
    saved_plan = tmp_path / "run-plan.json"
    plan.write(saved_plan)
    run = da.run(plan, out=tmp_path / "run")
    path.unlink()
    assert da.inspect(python_pool, verify=True).retained_parts == 2
    assert da.inspect(run, verify=True).accepted == 1
    assert read_source(saved_plan).plan_id == plan.plan_id


def test_excel_requires_unambiguous_sheet_and_rejects_mapped_formulas(tmp_path: Path):
    path = tmp_path / "parts.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "notes"
    sheet = book.create_sheet("parts")
    sheet.append(["part_id", "sequence", "ignored"])
    sheet.append(["p1", "ACGT", "=1+1"])
    book.save(path)
    with pytest.raises(ValueError, match="sheet"):
        read_parts(parts.PartTable(path, "xlsx"))
    source = parts.PartTable(path, "xlsx", sheet="parts")
    assert read_parts(source).parts[0].sequence == "ACGT"
    assert table_from_dict(table_to_dict(source)).sheet == "parts"
    sheet["B2"] = '="ACGT"'
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match=r"formula.*B2|B2.*formula"):
        read_parts(source)


@pytest.mark.parametrize("format_name", ["parquet", "xlsx"])
@pytest.mark.parametrize("invalid", [1.5, True])
def test_typed_coordinates_are_not_silently_coerced(
    tmp_path: Path, format_name: str, invalid: object
):
    path = tmp_path / f"parts.{format_name}"
    write_table(
        path,
        [
            {
                "part_id": "p1",
                "sequence": "ACGT",
                "core_start": invalid,
                "core_end": 3,
                "core_orientation": "forward",
            }
        ],
    )
    with pytest.raises(ValueError, match=r"row 1.*core_start"):
        read_parts(parts.PartTable(path, format_name))


@pytest.mark.parametrize("format_name", ["parquet", "xlsx"])
@pytest.mark.parametrize("sequence", [None, 123])
def test_required_text_is_not_inferred_from_typed_values(
    tmp_path: Path, format_name: str, sequence: object
):
    path = tmp_path / f"parts.{format_name}"
    write_table(path, [{"part_id": "p1", "sequence": sequence}])
    with pytest.raises(ValueError, match=r"row 1.*sequence"):
        read_parts(
            parts.PartTable(
                path, format_name, normalization=parts.Normalization(uppercase=True)
            )
        )


def test_sheet_option_is_specific_to_excel_and_csv_wire_is_unchanged():
    source = parts.PartTable("parts.csv", "csv")
    assert "sheet" not in table_to_dict(source)
    with pytest.raises(ValueError, match="sheet"):
        parts.PartTable("parts.csv", "csv", sheet="parts")


def test_excel_values_outside_header_are_not_silently_discarded(tmp_path: Path):
    path = tmp_path / "parts.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["part_id", "sequence"])
    book.active.append(["p1", "AAA", "unlabelled value"])
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="header"):
        read_parts(parts.PartTable(path, "xlsx"))


@pytest.mark.parametrize(
    "format_name,dependency", [("parquet", "pyarrow"), ("xlsx", "openpyxl")]
)
def test_missing_optional_reader_fails_before_output_creation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, format_name: str, dependency: str
):
    path = tmp_path / f"parts.{format_name}"
    write_table(path, [{"part_id": "p1", "sequence": "AAA"}])
    recipe = parts.PreparationSpec(parts.PartTable(path, format_name))
    saved = tmp_path / "request.json"
    da.export(recipe, view="request", out=saved)
    original_import = builtins.__import__

    def unavailable(name: str, *args: object, **kwargs: object) -> object:
        if name.split(".", maxsplit=1)[0] == dependency:
            raise ModuleNotFoundError(name)
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", unavailable)
    out = tmp_path / "pool"
    result = CliRunner().invoke(
        app, ["prepare", str(saved), "--out", str(out), "--json"]
    )
    assert result.exit_code == 2, result.output
    assert "install 'dense-arrays[tables]'" in json.loads(result.stdout)["message"]
    assert not out.exists()


@pytest.mark.parametrize("format_name", ["parquet", "xlsx"])
def test_corrupt_table_has_a_structured_cli_error(tmp_path: Path, format_name: str):
    path = tmp_path / f"parts.{format_name}"
    path.write_bytes(b"not a table")
    recipe = parts.PreparationSpec(parts.PartTable(path, format_name))
    saved = tmp_path / "request.json"
    da.export(recipe, view="request", out=saved)
    result = CliRunner().invoke(app, ["plan", str(saved), "--json"])
    assert result.exit_code == 2, result.output
    assert json.loads(result.stdout)["code"] == "invalid_input"
