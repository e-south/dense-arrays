"""
--------------------------------------------------------------------------------
Dense Arrays
dense-arrays/src/dense_arrays/parts/tables/excel.py

Optional values-only Excel input with explicit worksheet selection.

Module Author(s): Eric J. South
Maintainer(s): Eric J. South
--------------------------------------------------------------------------------
"""  # noqa: D205, D400 - structured module header

from __future__ import annotations

import io
from contextlib import closing, contextmanager
from typing import TYPE_CHECKING
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from . import TableRows

if TYPE_CHECKING:
    from collections.abc import Iterator

    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    from dense_arrays.parts.models import PartTable


def _select_sheet(book: Workbook, source: PartTable) -> ReadOnlyWorksheet:
    names = [sheet.title for sheet in book.worksheets]
    if source.sheet is None and len(names) != 1:
        msg = f"{source.table}: select a sheet by name; available: {names}"
        raise ValueError(msg)
    name = source.sheet if source.sheet is not None else names[0]
    if name not in names:
        msg = f"{source.table}: unknown sheet {name!r}; available: {names}"
        raise ValueError(msg)
    return book[name]


@contextmanager
def open_excel(source: PartTable, data: bytes) -> Iterator[TableRows]:
    """Read mapped cells as stored values; formulas never use stale caches."""
    try:
        from openpyxl import load_workbook  # noqa: PLC0415 - optional input dependency
    except ImportError as err:
        msg = "Excel input requires openpyxl; install 'dense-arrays[tables]'"
        raise ValueError(msg) from err
    try:
        with (
            io.BytesIO(data) as stream,
            closing(
                load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            ) as book,
        ):
            sheet = _select_sheet(book, source)
            name = sheet.title
            sheet.reset_dimensions()
            with closing(sheet.iter_rows()) as cells:
                header = tuple(cell.value for cell in next(cells, ()))

                def rows(columns: tuple[str, ...]) -> Iterator[dict[str, object]]:
                    positions = {column: header.index(column) for column in columns}
                    for row in cells:
                        if any(cell.value is not None for cell in row[len(header) :]):
                            msg = (
                                f"{source.table}: sheet {name!r}: "
                                "row has values beyond its header columns"
                            )
                            raise ValueError(msg)
                        if all(cell.value is None for cell in row):
                            continue
                        result = {}
                        for column, position in positions.items():
                            cell = row[position] if position < len(row) else None
                            if cell is not None and cell.data_type in {"f", "e"}:
                                kind = "formula" if cell.data_type == "f" else "error"
                                msg = (
                                    f"{source.table}: sheet {name!r}, "
                                    f"{cell.coordinate}: {kind} cells require "
                                    "explicit stored values"
                                )
                                raise ValueError(msg)
                            result[column] = cell.value if cell is not None else None
                        yield result

                yield TableRows(header, rows)
    except (BadZipFile, ParseError) as err:
        msg = f"{source.table}: invalid xlsx workbook: {err}"
        raise ValueError(msg) from err
